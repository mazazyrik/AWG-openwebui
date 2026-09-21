from __future__ import annotations

import json
import resource
import sys
from pathlib import Path


def _limits() -> None:
    if sys.platform != 'darwin':
        resource.setrlimit(resource.RLIMIT_AS, (512 * 1024 * 1024, resource.getrlimit(resource.RLIMIT_AS)[1]))
    resource.setrlimit(resource.RLIMIT_CPU, (30, resource.getrlimit(resource.RLIMIT_CPU)[1]))
    resource.setrlimit(resource.RLIMIT_FSIZE, (64 * 1024 * 1024, resource.getrlimit(resource.RLIMIT_FSIZE)[1]))
    if hasattr(resource, 'RLIMIT_NPROC'):
        _, hard = resource.getrlimit(resource.RLIMIT_NPROC)
        resource.setrlimit(resource.RLIMIT_NPROC, (min(16, hard), hard))


def _append(parts: list[str], value, length: int, limit: int) -> int:
    if value is None:
        return length
    text = str(value).strip()
    if not text:
        return length
    length += len(text) + 1
    if length > limit:
        raise ValueError('Artifact text exceeds the approved limit')
    parts.append(text)
    return length


def _extract(path: Path, extension: str, limit: int) -> str:
    parts: list[str] = []
    length = 0
    if extension == 'pdf':
        from pypdf import PdfReader

        reader = PdfReader(path, strict=True)
        if len(reader.pages) > 1000:
            raise ValueError('PDF page limit exceeded')
        for page in reader.pages:
            length = _append(parts, page.extract_text() or '', length, limit)
    elif extension == 'docx':
        from docx import Document

        document = Document(path)
        for paragraph in document.paragraphs:
            length = _append(parts, paragraph.text, length, limit)
        for table in document.tables:
            for row in table.rows:
                length = _append(parts, '\t'.join(cell.text for cell in row.cells), length, limit)
    elif extension == 'xlsx':
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        if len(workbook.worksheets) > 256:
            raise ValueError('Spreadsheet sheet limit exceeded')
        for sheet in workbook.worksheets:
            length = _append(parts, sheet.title, length, limit)
            for row in sheet.iter_rows(values_only=True):
                length = _append(parts, '\t'.join('' if value is None else str(value) for value in row), length, limit)
        workbook.close()
    elif extension == 'pptx':
        from pptx import Presentation

        presentation = Presentation(path)
        if len(presentation.slides) > 1000:
            raise ValueError('Presentation slide limit exceeded')
        for slide in presentation.slides:
            for shape in slide.shapes:
                if hasattr(shape, 'text'):
                    length = _append(parts, shape.text, length, limit)
    else:
        raise ValueError('Artifact type is not approved')
    return '\n'.join(parts)


def main() -> None:
    if len(sys.argv) != 4:
        raise SystemExit(2)
    _limits()
    path = Path(sys.argv[1]).resolve(strict=True)
    extension = sys.argv[2]
    limit = int(sys.argv[3])
    text = _extract(path, extension, limit)
    sys.stdout.write(json.dumps({'text': text}, ensure_ascii=False))


if __name__ == '__main__':
    main()
